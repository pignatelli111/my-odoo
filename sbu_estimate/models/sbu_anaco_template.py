# -*- coding: utf-8 -*-
"""Official SBU ANACO workbook: same columns and formulas as the client sheets, protected, with a
hidden template id / version that the import checks (files changed by third parties are refused)."""
import base64
import io
import re

from odoo import _, models
from odoo.exceptions import UserError

TEMPLATE_SHEET = 'SBU_TEMPLATE'
TEMPLATE_ID = 'SBU-ANACO'
TEMPLATE_VERSION = 1
HEADER_ROW, PARAM_ROW, FIRST_ROW, LAST_ROW, TOTAL_ROW = 4, 5, 12, 156, 157
SAL_HEADER_ROW, SAL_FIRST_ROW, SAL_LAST_ROW = 15, 16, 215
SAL_SHEET = 'Voci Contrattuali_SAL'

# Inputs (editable) per row. Every other column holds a formula.
INPUT_COLUMNS = ('B', 'C', 'D', 'F', 'H', 'J', 'T', 'X', 'AD', 'AF', 'AH', 'AJ', 'AL', 'AN', 'AP', 'AR', 'AT',
                 'AX', 'AZ', 'BV')
# Simple component columns: CAD input → TOT = CAD × Qt.
COMPONENT_INPUTS = (
    ('AD', 'AE', 'Pacchetto di coibentazione'), ('AF', 'AG', 'Kit avvolgimento'), ('AH', 'AI', 'Cassonetto'),
    ('AJ', 'AK', 'Automatismo'), ('AL', 'AM', 'Zanzariera'), ('AN', 'AO', 'Vetro'), ('AP', 'AQ', 'Pannello'),
    ('AR', 'AS', 'Controtelaio'), ('AT', 'AU', 'Trasformazione'), ('AX', 'AY', 'Posa lamiera'),
    ('AZ', 'BA', 'Nolo'),
)
# Row 5: (column, label shown in row 3, estimate field or default, is a percentage)
PARAMETERS = (
    ('K', 'Sc 1 serramento', None, False), ('L', 'Sc 2 serramento', None, False),
    ('M', 'Sc 3 serramento', None, False), ('P', 'Staffame €/ml', 'anaco_staffame_rate_ml', False),
    ('Q', 'Sc 1 staffame', None, False), ('R', 'Sc 2 staffame', None, False),
    ('U', 'Sc accessori', None, False), ('Y', 'Sc 1 oscuram.', None, False), ('Z', 'Sc 2 oscuram.', None, False),
    ('AA', 'Sc 3 oscuram.', None, False), ('AV', 'Posa €/m²', 'anaco_posa_rate_sqm', False),
    ('AW', 'Posa MQ minimo', 'anaco_posa_min_sqm', False), ('BE', 'Trasporto %', 'anaco_pct_transport', True),
    ('BG', 'Sviluppo tecnico e PM %', 'anaco_pct_tech_pm', True), ('BI', 'Presidio cantiere %', 'anaco_pct_site', True),
    ('BK', 'Extra %', 'anaco_pct_extra', True), ('BM', 'Costi industriali %', 'anaco_pct_industrial', True),
    ('BP', 'MOL %', 'anaco_pct_mol', True),
)
_MULTIPLIERS = {'K', 'L', 'M', 'Q', 'R', 'U', 'Y', 'Z', 'AA'}
# Row 157 sums (the same totals as the client sheets and SUMMARY COSTI).
TOTAL_COLUMNS = ('I', 'O', 'S', 'W', 'AC', 'AE', 'AG', 'AI', 'AK', 'AM', 'AO', 'AQ', 'AS', 'AU', 'AW', 'AY', 'BA',
                 'BC', 'BF', 'BH', 'BJ', 'BL', 'BO', 'BR', 'BT')


def _headers():
    """Row 4 header of every template column (the import compares these)."""
    headers = {
        'B': 'Pos.', 'C': 'Descrizione', 'D': 'B (mm)', 'F': 'H (mm)', 'G': 'MQ/Cad.', 'H': 'Qt.',
        'I': 'MQ totali', 'J': 'Serramento listino CAD', 'K': 'Sc 1', 'L': 'Sc 2', 'M': 'Sc 3',
        'N': 'Serramento scontato CAD', 'O': 'Serramento scontato TOT',
        'P': 'ST / LZ staffame CAD', 'Q': 'Staffame sc 1', 'R': 'Staffame scontato CAD', 'S': 'Staffame TOT',
        'T': 'Accessorio maniglie listino', 'U': 'Accessorio sc 1', 'V': 'Accessorio maniglie CAD',
        'W': 'Accessorio maniglie TOT', 'X': 'Oscuramento listino', 'Y': 'Oscuramento sc 1',
        'Z': 'Oscuramento sc 2', 'AA': 'Oscuramento sc 3', 'AB': 'Oscuramento scontato CAD',
        'AC': 'Oscuramento scontato TOT', 'AV': 'Smontaggio posa ser. e oscur. CAD',
        'AW': 'Smontaggio posa ser. e oscur. TOT', 'BB': 'Costo lavorato e posato CAD',
        'BC': 'Costo lavorato e posato TOT', 'BD': 'Costo lavorato e posato / MQ',
        'BE': 'Trasporto e imballo CAD', 'BF': 'Trasporto e imballo TOT', 'BG': 'Sviluppo tecnico e PM CAD',
        'BH': 'Sviluppo tecnico e PM TOT', 'BI': 'Presidio cantiere CAD', 'BJ': 'Presidio cantiere TOT',
        'BK': 'Extra CAD', 'BL': 'Extra TOT', 'BM': 'Costi industriali %', 'BN': 'Costi industriali CAD',
        'BO': 'Costi industriali TOT', 'BP': 'MOL %', 'BQ': 'MOL CAD', 'BR': 'MOL TOT',
        'BS': 'Prezzo cliente CAD', 'BT': 'Prezzo cliente TOT', 'BU': 'Prezzo cliente / MQ', 'BV': 'Note',
    }
    for cad, tot, label in COMPONENT_INPUTS:
        headers[cad] = f'{label} CAD'
        headers[tot] = f'{label} TOT'
    return headers


def _row_formulas(r):
    f = {
        'G': f'=D{r}*F{r}/1000000', 'I': f'=G{r}*H{r}',
        'K': f'=J{r}*$K$5', 'L': f'=K{r}*$L$5', 'M': f'=L{r}*$M$5', 'N': f'=M{r}', 'O': f'=N{r}*H{r}',
        'P': f'=2*(D{r}+F{r})/1000*$P$5', 'Q': f'=P{r}*$Q$5', 'R': f'=Q{r}*$R$5', 'S': f'=R{r}*H{r}',
        'U': f'=T{r}*$U$5', 'V': f'=U{r}', 'W': f'=V{r}*H{r}',
        'Y': f'=X{r}*$Y$5', 'Z': f'=Y{r}*$Z$5', 'AA': f'=Z{r}*$AA$5', 'AB': f'=AA{r}', 'AC': f'=AB{r}*H{r}',
        'AV': f'=MAX(G{r},$AW$5)*$AV$5', 'AW': f'=AV{r}*H{r}',
        'BB': f'=N{r}+R{r}+V{r}+AB{r}+AD{r}+AF{r}+AH{r}+AJ{r}+AL{r}+AN{r}+AP{r}+AR{r}+AT{r}+AV{r}+AX{r}+AZ{r}',
        'BC': f'=O{r}+S{r}+W{r}+AC{r}+AE{r}+AG{r}+AI{r}+AK{r}+AM{r}+AO{r}+AQ{r}+AS{r}+AU{r}+AW{r}+AY{r}+BA{r}',
        'BD': f'=IF(I{r}=0,0,BC{r}/I{r})',
        'BE': f'=BB{r}*$BE$5', 'BF': f'=BC{r}*$BE$5', 'BG': f'=BB{r}*$BG$5', 'BH': f'=BC{r}*$BG$5',
        'BI': f'=BB{r}*$BI$5', 'BJ': f'=BC{r}*$BI$5', 'BK': f'=BB{r}*$BK$5', 'BL': f'=BC{r}*$BK$5',
        'BM': '=$BM$5', 'BN': f'=BB{r}*BM{r}', 'BO': f'=BC{r}*BM{r}',
        'BP': '=$BP$5', 'BQ': f'=BB{r}*BP{r}', 'BR': f'=BC{r}*BP{r}',
        'BS': f'=BB{r}+BE{r}+BG{r}+BI{r}+BK{r}+BN{r}+BQ{r}',
        'BT': f'=BC{r}+BF{r}+BH{r}+BJ{r}+BL{r}+BO{r}+BR{r}', 'BU': f'=IF(I{r}=0,0,BT{r}/I{r})',
    }
    for cad, tot, _label in COMPONENT_INPUTS:
        f[tot] = f'={cad}{r}*H{r}'
    return f


def _compact(text):
    return re.sub(r'[^0-9A-Z€²%]', '', (text or '').upper())


def build_anaco_template(params, password):
    """Bytes of the protected ANACO workbook; ``params`` maps a row-5 column to its value."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    sh = wb.active
    sh.title = 'ANACO'
    bold, header_fill = Font(bold=True), PatternFill('solid', fgColor='DDE7F3')
    input_fill = PatternFill('solid', fgColor='FFF7D6')
    unlocked = Protection(locked=False)
    sh['C2'] = 'ANALISI COSTO PREVENTIVO — modello SBU'
    sh['C2'].font = Font(bold=True, size=14)
    for col, text in _headers().items():
        cell = sh[f'{col}{HEADER_ROW}']
        cell.value, cell.font, cell.fill = text, bold, header_fill
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    for col, label, _field, _pct in PARAMETERS:
        sh[f'{col}3'] = label
        cell = sh[f'{col}{PARAM_ROW}']
        cell.value = params.get(col, 1.0 if col in _MULTIPLIERS else 0.0)
        cell.protection, cell.fill = unlocked, input_fill
    numbers = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                             showErrorMessage=True, error='Inserire un numero ≥ 0')
    sh.add_data_validation(numbers)
    for r in range(FIRST_ROW, LAST_ROW + 1):
        for col, formula in _row_formulas(r).items():
            sh[f'{col}{r}'] = formula
        for col in INPUT_COLUMNS:
            cell = sh[f'{col}{r}']
            cell.protection, cell.fill = unlocked, input_fill
            if col not in ('B', 'C', 'BV'):
                numbers.add(cell)
    for col in TOTAL_COLUMNS:
        sh[f'{col}{TOTAL_ROW}'] = f'=SUM({col}{FIRST_ROW}:{col}{LAST_ROW})'
    sh[f'C{TOTAL_ROW}'] = 'TOTALE'
    sh.freeze_panes = f'D{FIRST_ROW}'
    sh.column_dimensions['C'].width = 45
    for col in _headers():
        if col != 'C':
            sh.column_dimensions[col].width = 13
    _protect(sh, password)

    sal = wb.create_sheet(SAL_SHEET)
    sal_headers = {2: 'ITEM', 3: 'DESCRIZIONE', 4: 'QTA', 5: 'TOT ML', 6: 'MQ', 7: 'TOT MQ', 8: 'PREZZO UNITARIO',
                   9: 'RITENUTA %'}
    sal_headers.update({98 + i: f'SAL-{i + 1}' for i in range(10)})
    for col, text in sal_headers.items():
        cell = sal.cell(SAL_HEADER_ROW, col, text)
        cell.font, cell.fill = bold, header_fill
    for r in range(SAL_FIRST_ROW, SAL_LAST_ROW + 1):
        for col in sal_headers:
            cell = sal.cell(r, col)
            cell.protection, cell.fill = unlocked, input_fill
    _protect(sal, password)

    marker = wb.create_sheet(TEMPLATE_SHEET)
    marker['A1'], marker['B1'] = TEMPLATE_ID, TEMPLATE_VERSION
    marker.sheet_state = 'veryHidden'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _protect(sheet, password):
    sheet.protection.sheet = True
    sheet.protection.password = password
    sheet.protection.formatColumns = False  # widths can be changed
    sheet.protection.formatRows = False
    sheet.protection.insertColumns = True  # columns cannot be inserted / deleted
    sheet.protection.deleteColumns = True
    sheet.protection.insertRows = True
    sheet.protection.deleteRows = True


def check_anaco_template(env, wb_vals):
    """None for a legacy workbook (no SBU marker); the version for a valid SBU template.

    Raises UserError when the marker is foreign / newer or when column headers were changed.
    """
    if TEMPLATE_SHEET not in wb_vals.sheetnames:
        return None
    marker = wb_vals[TEMPLATE_SHEET]
    template_id, version = (marker['A1'].value or ''), marker['B1'].value
    if template_id != TEMPLATE_ID:
        raise UserError(env._('This file is a "%(found)s" template, not an ANACO estimate template.', found=template_id))
    try:
        version = int(version)
    except (TypeError, ValueError):
        raise UserError(env._('The template version is missing or damaged: download a new ANACO template.')) from None
    if version > TEMPLATE_VERSION:
        raise UserError(env._('This ANACO template (version %(v)s) is newer than this Odoo (version %(odoo)s): '
                          'update Odoo first.', v=version, odoo=TEMPLATE_VERSION))
    if 'ANACO' not in wb_vals.sheetnames:
        raise UserError(env._('The sheet «ANACO» was removed or renamed in this template.'))
    sheet = wb_vals['ANACO']
    changed = []
    for col, expected in _headers().items():
        found = sheet[f'{col}{HEADER_ROW}'].value
        if _compact(str(found or '')) != _compact(expected):
            changed.append(env._('column %(col)s: "%(found)s" instead of "%(expected)s"',
                             col=col, found=found or '', expected=expected))
    if changed:
        raise UserError(env._(
            'The ANACO template was modified (columns moved, renamed or inserted), so it cannot be imported '
            'safely:\n%(changes)s\n\nDownload a new template from the estimate and copy the data into it.',
            changes='\n'.join(changed[:10]) + ('\n…' if len(changed) > 10 else ''),
        ))
    return version


class SbuEstimate(models.Model):
    _inherit = 'sbu.estimate'

    def _sbu_anaco_template_params(self):
        self.ensure_one()
        params = {}
        for col, _label, field, is_pct in PARAMETERS:
            if field:
                value = self[field] or 0.0
                params[col] = value / 100.0 if is_pct else value
        params['Q'] = self.anaco_staffame_multiplier or 1.0
        return params

    def action_download_anaco_template(self):
        self.ensure_one()
        password = self.env['ir.config_parameter'].sudo().get_param('sbu.template_password', 'sbu')
        data = build_anaco_template(self._sbu_anaco_template_params(), password)
        attachment = self.env['ir.attachment'].create({
            'name': f'ANACO_{self.name.replace("/", "-")}.xlsx',
            'datas': base64.b64encode(data),
            'res_model': self._name,
            'res_id': self.id,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'new',
        }
