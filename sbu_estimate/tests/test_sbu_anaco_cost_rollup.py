# -*- coding: utf-8 -*-
"""ANACO engine = Excel formulas (values from ANACO_SCHIAFFINO21_PERIMETRO_9_EURO_ML.xlsx, row 13)."""
import base64
import io

from odoo.tests import tagged
from odoo.tests.common import TransactionCase

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

from odoo.addons.sbu_estimate.wizards.sbu_estimate_anaco_import_wizard import (
    _anaco_posa_rule,
    _anaco_staffame_rule,
)

SCHIAFFINO_PARAMS = {
    'anaco_pct_transport': 4.0,
    'anaco_pct_tech_pm': 10.0,
    'anaco_pct_site': 0.0,
    'anaco_pct_extra': 1.0,
    'anaco_pct_mol': 55.0,
    'anaco_staffame_rate_ml': 9.0,
    'anaco_posa_rate_sqm': 60.0,
    'anaco_posa_min_sqm': 1.8,
}


@tagged('post_install', '-at_install')
class TestSbuAnacoCostRollup(TransactionCase):

    def setUp(self):
        super().setUp()
        partner = self.env['res.partner'].create({'name': 'Vitruvius test'})
        self.estimate = self.env['sbu.estimate'].create(dict(SCHIAFFINO_PARAMS, partner_id=partner.id))

    def _row13(self, **extra):
        vals = {
            'estimate_id': self.estimate.id,
            'description': 'Serramento PVC 500×800',
            'width_mm': 500, 'height_mm': 800, 'qty': 6,
            'price_serramento_cad': 134.59,
            'cost_coibentazione_cad': 15.0,
        }
        vals.update(extra)
        return self.env['sbu.estimate.line'].create(vals)

    def test_row_equals_excel(self):
        line = self._row13()
        self.assertAlmostEqual(line.cost_staffame_cad, 23.4, msg='R = 2*(D+F)/1000*9')
        self.assertAlmostEqual(line.price_smontaggio_cad, 108.0, msg='AV = MAX(G,1.8)*60')
        self.assertAlmostEqual(line.cost_mol_pct, 55.0, msg='BP from $BP$5')
        expected = {  # Excel row 13
            'sqm': 2.4,                              # I
            'cost_material_worked_cad': 280.99,      # BB
            'cost_material_worked_tot': 1685.94,     # BC
            'cost_per_sqm': 702.475,                 # BD
            'cost_trasporto_cad': 11.2396,           # BE
            'cost_trasporto_tot': 67.4376,           # BF
            'cost_tech_pm_cad': 28.099,              # BG
            'cost_tech_pm_tot': 168.594,             # BH
            'cost_cantiere_cad': 0.0,                # BI
            'cost_extra_cad': 2.8099,                # BK
            'cost_extra_tot': 16.8594,               # BL
            'cost_industrial_cad': 0.0,              # BN
            'cost_mol_amount_cad': 154.5445,         # BQ
            'cost_mol_amount_tot': 927.267,          # BR
            'price_total_cad': 477.683,              # BS
            'price_total_tot': 2866.098,             # BT
            'price_per_sqm': 1194.2075,              # BU
        }
        for fname, value in expected.items():
            self.assertAlmostEqual(line[fname], value, places=6, msg=fname)
        self.assertAlmostEqual(line.cost_total_cad, 280.99 + 11.2396 + 28.099 + 2.8099)
        self.assertAlmostEqual(line.margin_amount, 927.267, places=6, msg='margin = MOL')

    def test_dimensions_and_parameters_recompute(self):
        line = self._row13()
        line.width_mm = 1000  # 1000×800: perimeter 3.6 m, 0.8 m² → minimum 1.8 m²
        self.assertAlmostEqual(line.cost_staffame_cad, 32.4)
        self.assertAlmostEqual(line.price_smontaggio_cad, 108.0)
        line.height_mm = 2000  # 2 m² > minimum
        self.assertAlmostEqual(line.price_smontaggio_cad, 120.0)
        self.estimate.anaco_pct_transport = 5.0
        self.assertAlmostEqual(line.cost_trasporto_cad, line.cost_material_worked_cad * 0.05)
        self.estimate.anaco_staffame_rate_ml = 10.0
        self.assertAlmostEqual(line.cost_staffame_cad, 60.0)

    def test_explicit_values_win_over_rules(self):
        line = self._row13(cost_staffame_cad=0.0, price_smontaggio_cad=95.0)
        self.assertAlmostEqual(line.cost_staffame_cad, 0.0)
        self.assertAlmostEqual(line.price_smontaggio_cad, 95.0)

    def test_manual_price_and_commission(self):
        line = self._row13(price_anaco_bs_cad=500.0)
        self.assertAlmostEqual(line.price_total_cad, 500.0)
        self.assertAlmostEqual(line.anaco_price_cad, 477.683)
        line.commission_pct = 10.0
        self.assertAlmostEqual(line.price_total_cad, 450.0)

    def test_estimate_totals_and_summary(self):
        self._row13()
        self._row13(width_mm=350, height_mm=1500, qty=1, price_serramento_cad=108.72)  # Excel row 12
        estimate = self.estimate
        self.assertAlmostEqual(estimate.anaco_material_worked_total, 1685.94 + 265.02, places=2)  # monetary: cents
        self.assertAlmostEqual(estimate.total_price, 2866.098 + 450.534, places=2)  # monetary: cents
        self.assertAlmostEqual(estimate.anaco_installation_total, 648.0 + 108.0, places=2)  # monetary: cents
        self.assertAlmostEqual(estimate.anaco_mol_total, 927.267 + 145.761, places=2)  # monetary: cents
        html = str(estimate.anaco_summary_html)
        self.assertIn('Smontaggio posa', html)
        self.assertIn('MOL', html)

    def test_formula_rules(self):
        param = {'P': 9.0, 'AV': 132.5}.get
        self.assertEqual(_anaco_staffame_rule('=2*(D13+F13)/1000*9', param), 9.0)
        self.assertEqual(_anaco_staffame_rule('=2*(D13+F13)/1000*$P$5', param), 9.0)
        self.assertIsNone(_anaco_staffame_rule('=P5*2', param))
        self.assertEqual(_anaco_posa_rule('=MAX(G13,1.8)*60', param), (1.8, 60.0))
        self.assertEqual(_anaco_posa_rule('=$AV$5*G13', param), (0.0, 132.5))
        self.assertIsNone(_anaco_posa_rule('=AV12', param))

    def test_import_workbook_with_rules_and_overrides(self):
        """Workbook saved without computed values: P / AV come from the formulas' rules."""
        if not openpyxl:
            self.skipTest('openpyxl not installed')
        wb = openpyxl.Workbook()
        sh = wb.active
        sh.title = 'ANACO'
        sh.cell(4, 14, 'Serramento\nScontato\nCAD.')
        sh.cell(4, 48, 'Smontaggio\nPosa Ser e Oscur.\nCAD')
        for col, value in {11: 1, 12: 1, 13: 1, 17: 1, 18: 1, 57: 0.04, 59: 0.1, 61: 0, 63: 0.01, 65: 0, 68: 0.55}.items():
            sh.cell(5, col, value)
        for row, (pos, w, h, qty, n_price, bs) in enumerate([
            ('ser', 500, 800, 6, 134.59, '=BB12+BE12+BG12+BI12+BK12+BN12+BQ12'),
            ('ser', 350, 1500, 1, 108.72, 999.0),
        ], start=12):
            sh.cell(row, 2, pos)
            sh.cell(row, 3, f'Serramento {w}x{h}')
            sh.cell(row, 4, w)
            sh.cell(row, 6, h)
            sh.cell(row, 8, qty)
            sh.cell(row, 14, n_price)
            sh.cell(row, 16, f'=2*(D{row}+F{row})/1000*9')
            sh.cell(row, 30, 15)
            sh.cell(row, 48, f'=MAX(G{row},1.8)*60')
            sh.cell(row, 71, bs)
        buf = io.BytesIO()
        wb.save(buf)
        wizard = self.env['sbu.estimate.anaco.import.wizard'].create({
            'estimate_id': self.estimate.id,
            'data_file': base64.b64encode(buf.getvalue()),
            'data_filename': 'schiaffino_like.xlsx',
            'import_anaco': True,
            'import_sal': False,
            'import_distinta_bom': False,
        })
        wizard.action_import()
        estimate = self.estimate
        self.assertAlmostEqual(estimate.anaco_staffame_rate_ml, 9.0)
        self.assertAlmostEqual(estimate.anaco_posa_rate_sqm, 60.0)
        self.assertAlmostEqual(estimate.anaco_posa_min_sqm, 1.8)
        self.assertAlmostEqual(estimate.anaco_pct_mol, 55.0)
        first, second = estimate.line_ids.sorted('sequence')
        self.assertAlmostEqual(first.cost_material_worked_cad, 280.99)
        self.assertAlmostEqual(first.price_total_cad, 477.683)
        self.assertFalse(first.price_anaco_bs_cad, 'BS formula: price computed')
        self.assertAlmostEqual(second.price_total_cad, 999.0, msg='BS typed in Excel is kept')
        self.assertIn('Smontaggio Posa Ser e Oscur.', estimate._sbu_anaco_labels()['price_smontaggio_cad'])
